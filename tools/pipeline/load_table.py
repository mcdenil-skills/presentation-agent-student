import csv, io, os
from charset_normalizer import from_path
from tools.lexicon import AUTHOR_COLUMNS


def detect_csv_dialect(path):
    """Возвращает (encoding, delimiter). Русский Excel часто cp1251 + ';'."""
    best = from_path(path).best()
    enc = best.encoding if best else "utf-8"
    sample = str(best) if best else open(path, encoding=enc, errors="replace").read()
    head = sample.splitlines()[0] if sample.splitlines() else ""
    delim = ";" if head.count(";") >= head.count(",") and ";" in head else ","
    if "\t" in head and head.count("\t") > head.count(delim):
        delim = "\t"
    return enc, delim


def _norm_rows(records, text_col, sheet="csv"):
    out = []
    for i, rec in enumerate(records, 1):
        txt = (rec.get(text_col) or "").strip()
        if not txt:
            continue
        author = next((rec.get(c) for c in AUTHOR_COLUMNS if rec.get(c)), None)
        out.append({"лист": sheet, "строка": i, "автор": author, "текст": txt})
    return out


def load_csv(path, text_col):
    enc, delim = detect_csv_dialect(path)
    text = str(from_path(path).best())
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    if reader.fieldnames and text_col not in reader.fieldnames:
        raise ValueError(f"Столбец '{text_col}' не найден. Доступные столбцы: {reader.fieldnames}")
    return _norm_rows(reader, text_col, sheet="csv")


def load_xlsx(path, text_col):
    # Старый .xls (BIFF) openpyxl не читает - отвергаем понятно, а не молча мусором
    if path.lower().endswith(".xls"):
        raise ValueError("Старый формат .xls не поддержан. Пересохрани таблицу в .xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    all_cols, found_col = set(), False
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(c) if c is not None else "" for c in next(it)]
        except StopIteration:
            continue
        all_cols.update(h for h in header if h)
        if text_col not in header:
            continue
        found_col = True
        ti = header.index(text_col)
        ai = next((header.index(h) for h in AUTHOR_COLUMNS if h in header), None)
        for j, r in enumerate(it, 1):
            txt = (str(r[ti]).strip() if ti < len(r) and r[ti] is not None else "")
            if not txt:
                continue
            author = (str(r[ai]) if ai is not None and ai < len(r) and r[ai] is not None else None)
            rows.append({"лист": ws.title, "строка": j, "автор": author, "текст": txt})
    if not found_col:
        raise ValueError(
            f"Столбец '{text_col}' не найден ни на одном листе. Доступные столбцы: {sorted(all_cols)}")
    return rows


def preview_columns(path, n=3):
    """Структура источника для показа владельцу: листы, столбцы, первые n строк."""
    if path.lower().endswith(".xls"):
        raise ValueError("Старый формат .xls не поддержан. Пересохрани таблицу в .xlsx")
    if path.lower().endswith((".xlsx",)):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(c) for c in next(it)]
        sample = [list(map(lambda x: str(x) if x is not None else "", row)) for _, row in zip(range(n), it)]
        return {"sheets": [w.title for w in wb.worksheets], "columns": header, "sample_rows": sample}
    enc, delim = detect_csv_dialect(path)
    text = str(from_path(path).best())
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    header = next(reader)
    sample = [row for _, row in zip(range(n), reader)]
    return {"sheets": ["csv"], "columns": header, "sample_rows": sample}


def extract_sheet_id(url_or_id: str) -> str:
    """Достаёт ID таблицы из ссылки Google Sheets или возвращает строку как есть."""
    import re
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url_or_id)
    return m.group(1) if m else url_or_id.strip()


def load_google_sheet(url_or_id, text_col, token_path, tmp_dir="tools/work"):
    """Экспортирует Google Sheet в XLSX через Drive (scope drive) и читает ВСЕ листы.
    ВАЖНО: НЕ csv-export (он отдаёт только первый лист) - именно XLSX, чтобы не терять листы."""
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload
    creds = Credentials.from_authorized_user_file(token_path)
    drive = build("drive", "v3", credentials=creds)
    sid = extract_sheet_id(url_or_id)
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    req = drive.files().export_media(fileId=sid, mimeType=XLSX)
    os.makedirs(tmp_dir, exist_ok=True)
    out_path = os.path.join(tmp_dir, f"sheet_{sid}.xlsx")
    buf = io.FileIO(out_path, "wb")
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.close()
    return load_xlsx(out_path, text_col)
